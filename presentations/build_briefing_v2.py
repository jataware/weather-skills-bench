"""Concise briefing: native PowerPoint bullets and one dashboard screenshot."""
import json
import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.enum.text import PP_ALIGN
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'presentations'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='Weather agent benchmark — study design'
prs.core_properties.author='Weather Skills Benchmark'
prs.core_properties.subject='Briefing for Tess and Steve'
BLACK=RGBColor(17,17,17);GRAY=RGBColor(100,100,100)

def textbox(sl,x,y,w,h,content,size=22,bold=False,color=BLACK):
    sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
    tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color;p.space_after=Pt(4)
    return sh

def slide(title):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(255,255,255)
    textbox(s,.72,.62,11.95,.78,title,31,True)
    sh=textbox(s,12.15,7.09,.47,.22,str(len(prs.slides)),11,color=GRAY);sh.text_frame.paragraphs[0].alignment=PP_ALIGN.RIGHT
    return s

def bullets(s,items,y=1.70,size=23,spacing=22,h=4.7):
    sh=textbox(s,.85,y,11.65,h,'',size)
    tf=sh.text_frame
    for i,item in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.font.name='Arial';p.font.size=Pt(size);p.font.color.rgb=BLACK
        p.line_spacing=1.10;p.space_after=Pt(spacing)
        # DrawingML list properties must precede defRPr. Appending a bullet
        # after the font properties produces invalid ordering in PowerPoint.
        prop=p._p.get_or_add_pPr()
        prop.set('lvl','0');prop.set('marL',str(Inches(.30)))
        prop.set('indent',str(-Inches(.24)));prop.set('defTabSz',str(Inches(.30)))
        for tag,attrs in [('a:buSzPct',{'val':'100000'}),
                          ('a:buFont',{'typeface':'Arial'}),
                          ('a:buChar',{'char':'•'})]:
            el=OxmlElement(tag)
            for key,value in attrs.items():el.set(key,value)
            prop.insert_element_before(el,'a:tabLst','a:defRPr','a:extLst')
        tabs=OxmlElement('a:tabLst');stop=OxmlElement('a:tab')
        stop.set('pos',str(Inches(.30)));stop.set('algn','l');tabs.append(stop)
        prop.insert_element_before(tabs,'a:defRPr','a:extLst')
        if isinstance(item,tuple):
            r=p.add_run();r.text=item[0];r.font.bold=True
            r=p.add_run();r.text=item[1]
        else:p.text=item
    return sh

def notes(s,txt,sources):
    s.notes_slide.notes_text_frame.text=txt+'\n\nSources:\n'+'\n'.join(sources)

s=slide('Questions for the benchmark')
bullets(s,[
    'Does reusing vetted code improve accuracy or reduce the cost of a complete forecasting workflow?',
    'Our skills teach agents how to use library APIs from Python. weather-skills primarily exposes parameterized commands.',
    'Working hypothesis: agents need Python to recover from data, parameter, and workflow errors.',
    'Compare task success, tokens, cost, and completion time across model sizes. Do not assume skills will win.'
])
notes(s,'The audience already knows acmadDL and africaS2S, so library introductions are omitted. Their guides cover API conventions, data shapes, and workflow/statistical discipline; scientific implementations are in the libraries. weather-skills also has documentation, but the current harness exposes its operations through subprocess wrappers. The value of reuse and the need for Python recovery are hypotheses, not established conclusions.', ['User-provided framing','../../accord/acmadDL/skills/acmaddl/SKILL.md','../../accord/africas2s/skills/africas2s/SKILL.md'])

s=slide('Current experiment: three execution conditions')
bullets(s,[
    ('Python only: ', 'the agent writes, runs, and revises its own code using standard scientific Python packages.'),
    ('Skills only: ', 'the agent reads guides and invokes weather-skills operations; custom Python is disabled.'),
    ('Python + skills: ', 'the agent can combine catalog operations with its own Python. Skill use is measured, not assumed.'),
    'Ten models, one task, three conditions: 30 planned attempts. The small-model tier includes 1B, 3B, 7B, 8B, and 9B models.'
],spacing=21)
textbox(s,.86,6.43,11.55,.49,'Astra is excluded because it helped develop the benchmark. Earlier runs remain archived.',16,color=GRAY)
notes(s,'The main comparison is Python only versus Python + skills; Skills only diagnoses the completeness and usability of the command interface. All conditions can receive execution feedback and retry within equal budgets. This is a weather-skills pilot, not yet a direct library evaluation. One repetition per task/model/condition is planned; no broad model ranking follows from it. Llama 1B/3B endpoints lack enforced JSON output; their prompts, parsing, and feedback stay the same while the unsupported response_format field is omitted. The ten models are Claude Fable 5.1, Claude Sonnet 5.5, Gemini 3.1 Flash-Lite, DeepSeek V4.1 Flash, Qwen3.5 9B, Ministral 3B, Llama 3.2 1B/3B, Qwen2.5 7B, and Llama 3.1 8B.', ['configs/refined-heat-panel-v4.json','configs/refined-heat-small-v4.json','METHODOLOGY.md'])

s=slide('Current task: a two-week Kenya heat outlook')
bullets(s,[
    'Use a real, archived ECMWF S2S forecast to calculate the hottest daily regional-mean temperature in each week.',
    'Required order: cosine-latitude spatial mean for each member/day → weekly peak per member → ensemble median, spread, minimum, and maximum.',
    'Pass requires a structured answer matching an independent numeric reference, correct units and source citation, and a valid, nonblank figure.',
    'Every attempt gets the same cached raw data and an empty workspace. Intermediate results stay within that attempt; no outputs or history cross runs.'
],size=22,spacing=20)
textbox(s,.86,6.36,11.6,.57,'Equal ceilings: 24 model calls, 40 executions, 200k tokens, 600 seconds.\nDownload and setup are excluded so solve time measures analysis, not data transfer.',16,color=GRAY)
notes(s,'Task: e2e-kenya-heat-cached-v3, executed under repaired protocol cached-heat-v4. Source issue 2026-09-27; week 1 lead days 1–7 and week 2 days 8–14. Service rectangle 5N/34E/5S/42E, boundary grid centres included; not a country-polygon average. All ensemble members including control receive equal weight. Sample spread uses ddof=1. The quantity is daily 2 m temperature, not a daily-maximum-temperature variable. Numerical absolute tolerance 1e-4, relative 1e-6; exact schema, units, and source. PNG validity, size, and nonblank pixels are checked; semantic visual correctness still needs review. This measures correct analysis of an issued forecast, not predictive skill against observations. Cached inputs are raw and hash-verified, not precomputed skill outputs. Successful unchanged skill calls may reuse their own artifacts within a run. Cold acquisition performance needs a separate test.', ['cases/e2e-kenya-heat-cached-v3.json','scripts/validate_refined.py','weather_bench/refined.py','weather_bench/execution_cache.py'])

# The landing page is the visual on this slide; no added diagram or duplicate heading.
s=prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(255,255,255)
path=OUT/'assets/dashboard-landing.png'
s.shapes.add_picture(str(path),0,0,width=prs.slide_width,height=Inches(7.0))
prov=json.loads((OUT/'assets/landing-provenance.json').read_text())
state='Completed study' if prov['finished'] and prov['recorded_attempts']==prov['planned_attempts'] else 'Partial study' if prov['finished'] else 'Interim snapshot'
caption=f"{state}: {prov['recorded_attempts']}/{prov['planned_attempts']} attempts. Select a result to inspect its trace."
textbox(s,.85,7.12,11.7,.25,caption,14,color=GRAY)
notes(s,'Actual unfiltered landing page captured at '+prov['exported_at']+'. '+str(prov['recorded_attempts'])+' of '+str(prov['planned_attempts'])+' attempts recorded. This is a viewport screenshot; the remaining model rows continue below the visible area. The dashboard supports success/time/token/cost charts and drilldowns to code, skill calls, model responses, errors, and grading. Provider errors are separated from task failures. The screenshot is an interface illustration, not a ranking.', ['docs/index.html','presentations/assets/landing-provenance.json'])

s=slide('Observed failure modes')
bullets(s,[
    ('Skills bypassed: ', '0 of 6 Python + skills attempts invoked a skill. The prompt makes skills optional; this condition tests availability, not actual reuse.'),
    ('Interface friction: ', 'wrong action names, XML tool-call wrappers, guide-read requirements, and invalid CLI arguments consumed turns.'),
    ('Recovery and delivery: ', 'Python indexing and unit errors required retries. Sonnet produced a figure but failed the strict artifact-submission contract.'),
    ('Budget exhaustion: ', '6 of 7 failures hit the token limit; Qwen’s hybrid run hit the time limit. Repeated context and unusable responses contributed.'),
],size=22,spacing=20)
textbox(s,.86,6.42,11.6,.42,'Initial six-model batch: 18 attempts, 7 failures. One attempt per model and condition.',16,color=GRAY)
notes(s,'Evidence is frozen to the completed initial six-model batch, not the still-expanding ten-model panel. All six hybrid attempts used Python and no skill invocations; only Ministral read one guide. Four hybrid attempts passed. The prompt begins with the Python agent instructions, presents skills as additional optional actions, and explicitly permits Python custom calculations. This supports an adoption/design diagnosis; it does not reveal a model’s internal motivation or establish a causal effect from prompt length. Skills-only Sonnet had 12 successful skill calls and a valid figure but rejected XML-like action wrappers and a literal [7,14] array under the artifact-only numeric submission rule. Its intermediate scientific correctness was not independently established here. Qwen strict-skills had nine protocol errors, including two empty responses ending at the output limit; its clipping and conversion worked. Gemini and Ministral strict-skills also encountered CLI arguments and dimension/precondition mismatches. Qwen hybrid summed six longitude cells without averaging, then timed out. Six failed attempts ended at the cumulative token ceiling, one at the shared task deadline. The follow-up should separately test a skill-discovery/preference policy with Python recovery, fix action transport and submission friction, and repeat paired runs. Existing records and scores remain unchanged.', ['results/initial-panel-failure-modes.json','results/qwen35-heat-failure-analysis.json','weather_bench/runner.py: BASE_PROMPT and SKILL_PROMPT'])

s=slide('Next: test the libraries and their skill guides directly')
bullets(s,[
    ('Baseline: ', 'general scientific Python; the agent implements the workflow itself.'),
    ('Libraries: ', 'the same Python environment plus acmadDL and africaS2S, with ordinary API documentation.'),
    ('Libraries + skills: ', 'the same libraries plus agent guides and progressively disclosed workflow references.'),
    'Keep Python available in every condition. Use seasonal tasks covering data acquisition, downscaling, held-out verification, and forecast delivery.'
],size=23,spacing=23)
textbox(s,.86,6.37,11.55,.54,'Predeclare tasks and scoring; freeze versions and seeds; repeat across regions, methods, and models.',17,color=GRAY)
notes(s,'This proposed library ablation has not been run in the current benchmark. Comparing baseline with libraries estimates reusable-code value; comparing libraries with libraries plus skills estimates additional guidance value. All arms retain Python and recovery. No condition should be deprived of basic API usability; ordinary documentation is the middle-arm control. Seasonal tasks should avoid training/verification leakage, use fixed or appropriately nested method selection, common valid masks, and independent references or cross-implementation checks. Agent correctness and meteorological predictive skill are different: held-out forecast verification may use RPSS, reliability, RMSE, or other predeclared metrics. Separate cold acquisition/setup from warm-data solve time. Use repeated tasks and model sampling before drawing claims about generalization.', ['User-provided objective','../../accord/acmadDL/skills/acmaddl/SKILL.md','../../accord/africas2s/skills/africas2s/SKILL.md','STATISTICS.md'])

suffix=os.environ.get('BRIEFING_SUFFIX','v2')
path=OUT/f'weather-agent-benchmark-briefing-{suffix}.pptx';prs.save(path)
print(path)
print('Slides:',len(prs.slides),'Native bullets:',sum(len(sh._element.xpath('.//a:buChar')) for s in prs.slides for sh in s.shapes),'Screenshots:',sum(sh.shape_type==13 for s in prs.slides for sh in s.shapes))
