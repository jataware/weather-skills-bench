"""Build editable PowerPoint slides. Requires python-pptx and Pillow.
Run: PYTHONPATH=/tmp/weather-slides-deps .venv/bin/python presentations/build_briefing.py
Only dashboard screenshots are raster; diagrams, tables, and chart are native.
"""
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'presentations'; ASSETS=OUT/'assets'
snap=json.loads((ASSETS/'study-snapshot.json').read_text())
s=snap['study']; rs=s['runs']; n=len(rs); planned=s['planned_runs']
prov=json.loads((ASSETS/'screenshot-provenance.json').read_text())
llm=sum(r['llm_seconds'] for r in rs); exe=sum(r['execution_seconds'] for r in rs); solve=sum(r['solve_seconds'] for r in rs)
share=round(100*llm/solve) if solve else 0
stamp=snap['exported_at'][:16].replace('T',' ')+' UTC'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='Benchmarking agents for weather forecasting'
prs.core_properties.subject='Reusable libraries, agent skills, and deterministic evaluation'
prs.core_properties.author='Weather Skills Benchmark'
prs.core_properties.comments='Editable PowerPoint. Native shapes, tables and chart. Dashboard screenshots are embedded raster captures. Interim data is explicitly labeled.'
INK='111111';GRAY='666666';RULE='D1D1D1';PALE='F2F2F2';BLUE='2057A5';WHITE='FFFFFF'

def rgb(c):return RGBColor.from_string(c)
def text(sl,x,y,w,h,string,size=20,bold=False,color=INK,align=None):
    sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(string.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(2)
        if align is not None:p.alignment=align
    return sh

def line(sl,x1,y1,x2,y2,color=RULE,width=1):
    sh=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(width);return sh

def rect(sl,x,y,w,h,fill=WHITE,stroke=INK):
    sh=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill);sh.line.color.rgb=rgb(stroke);sh.line.width=Pt(1);sh._element.spPr.append(OxmlElement("a:effectLst"));return sh

def arrow(sl,x,y,w=.38,h=.16,color=INK):
    sh=sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(color);sh.line.fill.background();return sh

def box(sl,x,y,w,h,title,body='',accent=False):
    rect(sl,x,y,w,h,WHITE,BLUE if accent else INK)
    text(sl,x+.18,y+.15,w-.36,.43,title,18,True,BLUE if accent else INK)
    if body:text(sl,x+.18,y+.65,w-.36,h-.75,body,17)

def header(title,kicker='WEATHER AGENT BENCHMARK'):
    sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=rgb(WHITE)
    text(sl,.58,.30,11,.22,kicker,10,True,GRAY)
    text(sl,.58,.81,12.1,.82,title,30,True)
    line(sl,.58,6.98,12.75,6.98)
    text(sl,.58,7.10,11.4,.2,'Weather Skills Benchmark  /  Research design and current pilot',10,color=GRAY)
    text(sl,12.2,7.08,.55,.23,f'{len(prs.slides):02d}',10,color=GRAY,align=PP_ALIGN.RIGHT)
    return sl

def note(sl,body,sources):
    sl.notes_slide.notes_text_frame.text=body+'\n\nSources / audit trail:\n'+'\n'.join(sources)

def table(sl,x,y,width,heights,colwidths,rows,fontsize=18,highlight=None):
    sh=sl.shapes.add_table(len(rows),len(colwidths),Inches(x),Inches(y),Inches(width),Inches(sum(heights)))
    tb=sh.table
    for j,cw in enumerate(colwidths):tb.columns[j].width=Inches(cw)
    for i,row in enumerate(rows):
        tb.rows[i].height=Inches(heights[i])
        for j,val in enumerate(row):
            c=tb.cell(i,j);c.text=val;c.margin_left=Inches(.13);c.margin_right=Inches(.13);c.margin_top=Inches(.12);c.margin_bottom=Inches(.08);c.vertical_anchor=MSO_ANCHOR.MIDDLE
            c.fill.solid();c.fill.fore_color.rgb=rgb(INK if i==0 else PALE if i%2 else WHITE)
            for p in c.text_frame.paragraphs:
                p.font.name='Arial';p.font.size=Pt(fontsize);p.font.bold=i==0 or j==0;p.font.color.rgb=rgb(WHITE if i==0 else BLUE if highlight==i and j==0 else INK)
    return sh

def picture(sl,path,x,y,w,h):
    iw,ih=Image.open(path).size;scale=min(w/iw,h/ih);ww=iw*scale;hh=ih*scale
    return sl.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))

# 1 — Purpose and actual library responsibilities.
a=header('Benchmarking agents for weather forecasting')
text(a,.6,1.75,11.9,.9,'Which combination of model, libraries, and agent guidance delivers correct forecasts with the least time, tokens, and cost?',25)
box(a,.62,3.15,3.70,2.1,'acmadDL','Fetch and normalize weather / climate data.\nReturn consistent xarray datasets.')
arrow(a,4.48,4.08,.43,.2)
box(a,5.08,3.15,4.03,2.1,'africaS2S','Downscale, calibrate, combine, and verify seasonal forecasts.')
arrow(a,9.29,4.08,.43,.2)
box(a,9.9,3.15,2.8,2.1,'Forecast products','Fields, probabilities, verification, and briefing figures.')
text(a,.65,5.66,12,.45,'Methods include logistic regression, canonical correlation analysis, quantile mapping, and BCSD.',18,color=GRAY)
text(a,.65,6.29,12,.38,'The benchmark tests the value of reuse against an agent implementing the workflow in Python.',19,True)
note(a,'The aim is an empirical comparison, not a presumption that reusable code always wins. BCSD means bias correction and spatial disaggregation. The user-facing library names are retained here; local documentation also uses Rosetta/deepscale aliases. The current pilot is not yet a direct test of acmadDL or africaS2S.', ['User-provided project context','../../accord/acmadDL/skills/acmaddl/SKILL.md','../../accord/africas2s/skills/africas2s/SKILL.md'])

# 2 — Interfaces, not different meanings of scientific correctness.
a=header('Skills can expose commands or teach library APIs')
text(a,.65,1.73,12,.4,'Both approaches reuse scientific code. They differ in how the agent composes and repairs the workflow.',20)
text(a,.65,2.49,2.45,.75,'Command-oriented\nweather-skills',18,True)
box(a,3.06,2.35,2.55,1.37,'Agent','Select skill and args')
arrow(a,5.77,2.95,.36,.18)
box(a,6.29,2.35,2.63,1.37,'Skill wrapper','Run a subprocess')
arrow(a,9.08,2.95,.36,.18)
box(a,9.59,2.35,3.05,1.37,'Vetted code','Write an artifact')
line(a,.65,4.02,12.67,4.02)
text(a,.65,4.43,2.45,.85,'Library-oriented\nacmadDL / africaS2S',19,True)
box(a,3.06,4.28,2.55,1.53,'Skill guide','API rules and\nrelevant references')
arrow(a,5.77,4.96,.36,.18)
box(a,6.29,4.28,2.63,1.53,'Python code','Compose library calls;\nhandle errors',True)
arrow(a,9.08,4.96,.36,.18)
box(a,9.59,4.28,3.05,1.53,'Vetted libraries','Execute scientific\nmethods')
text(a,.65,6.22,12,.51,'Hypothesis: progressive disclosure plus Python supports recovery from data, API, and workflow failures.',20,True)
note(a,'The library skills are readme-like guides with deeper references loaded when needed. They teach units, dimensions, API use, and statistical discipline. Numerical implementations are maintained in the libraries; generated Python handles composition and exceptions. weather-skills also has guides: this is a distinction in the execution interface, not a claim that its guidance is absent. The hypothesis that general code execution improves recovery remains to be evaluated quantitatively.', ['weather-skills-catalog skill scripts and guides','../../accord/acmadDL/skills/acmaddl/SKILL.md','../../accord/africas2s/skills/africas2s/SKILL.md'])

# 3 — Current intervention.
a=header('The current pilot compares three execution conditions')
table(a,.65,1.84,12.03,[.57,.76,.76,.76],[2.47,2.32,3.37,3.87],[['Condition','Write Python','Use weather-skills','Question'],['Python only','Yes','No','Can the agent implement the task?'],['Skills only','No','Required','Can catalog operations complete it?'],['Python + skills','Yes','Available','Does reuse help when code is allowed?']],17,3)
text(a,.65,4.96,12,.43,'Main comparison: Python only versus Python + skills. Skills only is a diagnostic condition.',20,True)
text(a,.65,5.57,12,.61,'10 models × 3 conditions × 1 task × 1 repetition = 30 planned attempts.\nThe panel spans frontier models and open-weight models from 1B to 9B.',18)
text(a,.65,6.41,12,.33,'Astra is excluded because it helped develop the benchmark. Earlier runs remain archived.',15,color=GRAY)
note(a,'Current models: Claude Fable 5.1, Claude Sonnet 5.5, Gemini 3.1 Flash-Lite, DeepSeek V4.1 Flash, Qwen3.5 9B, Ministral 3B, Llama 3.2 1B, Llama 3.2 3B, Qwen2.5 7B, Llama 3.1 8B. DeepSeek is a large open-weight model; the 1B–9B statement describes the smaller tier, not every open model. All arms receive execution feedback and may retry. No one-shot baseline is included. Skills are available, not mandatory, in Python + skills; invocation is measured. Llama 1B/3B endpoints lack API-enforced JSON mode but receive the same action instructions and parsing. Astra exclusion reduces an authoring conflict; it does not eliminate all design bias. Current comparisons concern weather-skills, not the two project libraries.', ['configs/refined-heat-panel-v4.json','configs/refined-heat-small-v4.json','METHODOLOGY.md'])

# 4 — Concrete task and ground truth.
a=header('Current task: a two-week heat outlook for Kenya')
text(a,.65,1.73,12,.68,'Use a frozen ECMWF S2S forecast to summarize the hottest daily regional-mean 2 m temperature in each week, including ensemble uncertainty.',22)
labels=[('1  Select','Service rectangle\n5°N–5°S, 34°E–42°E'),('2  Spatial mean','Cosine-latitude weights\nPer member and day'),('3  Weekly peak','Days 1–7 and 8–14\nFor each member'),('4  Ensemble','Median, sample SD,\nminimum, maximum')]
for i,(title,body) in enumerate(labels):
    x=.65+i*3.04;box(a,x,2.8,2.79,1.76,title,body)
    if i<3:arrow(a,x+2.84,3.61,.16,.12)
text(a,.65,4.88,12,.43,'Required order: spatial mean → weekly peak per member → ensemble statistics.',20,True)
text(a,.65,5.51,6,.88,'Deliverables\nStructured answer, source citation, and a labeled uncertainty figure.',18)
text(a,7.03,5.51,5.62,.94,'Deterministic checks\nIndependent numeric reference; exact schema and units; valid, nonblank PNG.',18)
text(a,.65,6.59,12,.2,'Forecast issue: 2026-09-27. This tests analysis of an issued forecast, not its predictive skill against observations.',12,color=GRAY)
note(a,'The task uses real, unmodified archived forecast bytes, not synthetic weather observations. Week 1 covers September 28–October 4; week 2 October 5–11. The service rectangle is not a Kenya polygon. All perturbed members and the control receive equal member weights. The quantity is daily 2 m temperature, not a daily-maximum-temperature variable. Numerical tolerance is abs=1e-4, rel=1e-6, with exact schema, dates/lead values, units, and source citation. The oracle uses an independent Python calculation; actual catalog workflows are also validated. A valid nonblank PNG is required, but chart semantics and aesthetics are not fully automatically graded. Workflow conformance is separate from answer correctness.', ['cases/e2e-kenya-heat-cached-v3.json','scripts/validate_refined.py','results/refined-reference-v4.json','https://storage.googleapis.com/kenya-forecasting-data/2026-09-27/data/ECMWF_s2s_daily_vars_2026-09-27.zarr'])

# 5 — Efficiency, fairness, isolation.
a=header('Cache the inputs; keep each attempt independent')
box(a,.65,2.10,3.25,2.26,'Raw forecast archive','Real raw forecast bytes\nVerified by content hashes\nMounted read-only')
for y in [2.0,3.03,4.06]:
    line(a,4.08,3.23,4.48,3.23,INK)
    line(a,4.48,3.23,4.48,y+.38,INK)
    arrow(a,4.51,y+.30,.4,.17)
for y,title in [(2.0,'Python only'),(3.03,'Skills only'),(4.06,'Python + skills')]:
    rect(a,5.08,y,7.57,.78)
    text(a,5.26,y+.21,2.45,.37,title,18,True)
    text(a,8.0,y+.20,4.38,.42,'Fresh conversation + empty workspace',17)
text(a,.65,5.13,5.72,.80,'Within a run\nReuse unchanged successful skill outputs; retain Python artifacts.',18)
text(a,7.04,5.13,5.61,.80,'Between runs\nShare raw inputs only. No generated outputs, intermediate results, or history.',18)
text(a,.65,6.29,12,.41,'Equal ceilings: 24 model calls  ·  40 executions  ·  200k tokens  ·  600 seconds',19,True)
note(a,'This is a warm-data benchmark. Download and environment setup are excluded from solve time; this improves repeatability and removes network transfer as a confound. The experiment does not measure live acquisition reliability or cold-start latency. Every run remains isolated. The raw cache does not contain transformed answers. Successful identical skill calls reuse artifacts only within the same run and only if inputs and outputs remain unchanged; cache hits are traced. Python agents can keep their own intermediate files. The repaired v4 catalog pins additional temperature/rolling/median fixes. Task IDs retain v3 naming, but protocol_version is cached-heat-v4. A full live-acquisition benchmark is a separate next-stage measurement.', ['weather_bench/refined.py','weather_bench/execution_cache.py','configs/refined-heat-panel-v4.json','results/refined-reference-v4.json'])

# 6 — Measurement and honest interim timing.
a=header('Measure correctness, resource use, and failure causes')
rows=[('Task success','Numeric reference + output contract'),('Workflow','Guide reads, calls, dependency order'),('Efficiency','Tokens, cost, solve time; cost per success'),('Diagnostics','API errors, invalid actions, execution errors,\nbudget exhaustion, wrong results')]
for i,(title,body) in enumerate(rows):
    y=1.95+i*1.0;line(a,.65,y-.12,7.02,y-.12)
    text(a,.65,y,1.70,.62,title,18,True);text(a,2.51,y,4.44,.73,body,18)
text(a,7.64,1.92,4.9,.43,'Where the elapsed time goes',20,True)
text(a,7.64,2.57,4.8,.94,f'{share}%',49,True,BLUE)
text(a,7.64,3.47,4.87,.57,'of recorded solve time was spent\nwaiting for model / API responses.',18)
cd=CategoryChartData();cd.categories=['Recorded time']
for name,value in [('Model / API',llm),('Execution',exe),('Other',max(0,solve-llm-exe))]:cd.add_series(name,[100*value/solve if solve else 0])
chart=a.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED,Inches(7.6),Inches(4.35),Inches(4.9),Inches(1.02),cd).chart
chart.has_legend=False;chart.has_title=False
chart.category_axis.visible=False;chart.value_axis.visible=False;chart.value_axis.has_major_gridlines=False;chart.value_axis.minimum_scale=0;chart.value_axis.maximum_scale=100
chart.plots[0].gap_width=30
for series,color in zip(chart.series,[BLUE,'777777','DDDDDD']):series.format.fill.solid();series.format.fill.fore_color.rgb=rgb(color);series.format.line.fill.background()
for x,color,label in [(7.65,BLUE,'Model / API'),(9.4,'777777','Execution'),(11.0,'DDDDDD','Other')]:
    rect(a,x,5.36,.12,.12,color,color);text(a,x+.2,5.28,1.45,.35,label,12,color=GRAY)
text(a,7.64,5.89,4.9,.57,f'Interim snapshot: {n}/{planned} attempts.\nTiming diagnosis, not a model ranking.',15,color=GRAY)
text(a,.65,6.39,6.9,.34,'Report service availability separately from task capability.',17,True)
note(a,f'The timing chart is native PowerPoint and uses summed run timing, not an unweighted mean of percentages. Snapshot {stamp}; {n} recorded of {planned} planned. Model/API time {llm:.3f}s; execution {exe:.3f}s; solve {solve:.3f}s. Provider queueing, inference, network response time and request failures are not separately identifiable in this aggregate. Default capability charts exclude terminal provider errors; operational results and reported spending retain them. Task timeouts and malformed responses remain task outcomes. Cost per successful task must include the stated attempt population and flag incomplete billing. Total tokens include repeated context and provider-cached input. Exact McNemar comparisons and Wilson intervals are available, but one task and one repetition do not support broad significance or a model ranking. Benchmark sampling and repetitions are still needed.', ['presentations/assets/study-snapshot.json','weather_bench/report.py','STATISTICS.md'])

# 7 — Dashboard comparison screenshot.
a=header('Inspect each comparison in the dashboard')
text(a,.65,1.72,12,.54,'Filter by experiment, model, and task; switch between success, time, tokens, and cost.',21)
picture(a,ASSETS/'dashboard-comparison.png',.62,2.47,12.08,3.65)
text(a,.65,6.31,12,.43,'Select a bar to inspect its attempts. Screenshot shows one model as an interface example.',18)
note(a,f'Actual dashboard screenshot, frozen at {stamp}. Selected model is Claude Fable 5.1 and chart metric is completion time. This is an interface demonstration from a partially completed study, not a claim about relative model performance. Missing attempts remain pending. Image captures genuine DOM content and original series colors. All surrounding slide elements are native PowerPoint.', ['docs/index.html','presentations/assets/screenshot-provenance.json','presentations/assets/dashboard-comparison.png'])

# 8 — Trace screenshot and auditing.
a=header('Trace a result back to the model and its executions')
picture(a,ASSETS/'dashboard-trace.png',.62,1.81,7.70,4.86)
line(a,8.68,1.96,8.68,6.56)
text(a,9.02,1.98,3.66,.52,'Each attempt records',21,True)
for y,title,body in [(2.87,'Model calls','Visible responses, request time,\ntokens, and billed cost.'),(4.0,'Execution log','Python or skill arguments,\noutputs, errors, and retries.'),(5.13,'Answer & grading','Submitted artifacts, reference\nvalues, and failed checks.')]:
    text(a,9.02,y,3.63,.35,title,19,True);text(a,9.02,y+.44,3.63,.7,body,17)
note(a,f'Actual run-detail drawer for run {prov["run_id"]}, with Model calls selected. The dashboard also supports execution traces and answer/grade inspection. Logs are sanitized; they exclude credentials and hidden reasoning. Reading a guide and invoking a skill are separately recorded. Correct alternative Python implementations may fail reference-recipe conformance but still pass scientific correctness. The screenshot is an audit example, not a sample chosen for aggregate inference.', ['presentations/assets/screenshot-provenance.json','presentations/assets/dashboard-trace.png','weather_bench/report.py: public_audit','docs/app.js: showRun'])

# 9 — Direct future test of project libraries.
a=header('Next: isolate the value of the libraries and their skills')
text(a,.65,1.73,12,.50,'The current pilot tests execution patterns. A direct acmadDL / africaS2S study needs its own tasks.',21)
table(a,.65,2.48,12.03,[.58,.71,.71,.71],[3.23,4.28,4.52],[['Proposed condition','What the agent receives','What the comparison estimates'],['Python baseline','General scientific Python','Reference for custom implementations'],['Python + libraries','Libraries + ordinary API documentation','Value of reusable implementations'],['Python + libraries + skills','Same libraries + agent skill guides','Added value of progressive disclosure']],17,3)
text(a,.65,5.55,12,.71,'Seasonal task: acquire hindcasts + observations → align data → fit / downscale → verify on held-out years → produce forecast probabilities and a briefing.',20,True)
text(a,.65,6.45,12,.30,'Predeclare tasks and scoring; freeze versions and seeds; repeat across regions, methods, and models.',17,color=GRAY)
note(a,'This proposed ablation is not implemented or run in the current study. Keep Python execution and recovery available in all three conditions. For the library-only arm, provide ordinary API documentation rather than withholding basic usability information; the skills arm adds workflow guidance and progressive disclosure. Fix candidate methods and evaluation folds or use properly nested selection to avoid optimistic verification. Reference checks should combine independently computed quantities, cross-implementation agreement, and invariants rather than grade a library exclusively against itself. Separate agent task correctness from meteorological predictive skill. Forecast skill needs held-out years, no calibration/tercile leakage, common masks, and prespecified metrics such as RPSS, reliability, or RMSE as appropriate. Include representative missing data, calendar differences, grids, and recoverable source errors. Report warm-data solve time and cold acquisition/setup separately. The decision is empirical: which configuration meets accuracy requirements at the lowest cost and latency, without assuming skills must win.', ['User-provided project objective','../../accord/acmadDL/skills/acmaddl/SKILL.md','../../accord/africas2s/skills/africas2s/SKILL.md','METHODOLOGY.md','STATISTICS.md'])

for sl in prs.slides:
    for sh in sl.shapes:
        for effect in sh._element.xpath('.//a:effectRef'):effect.set('idx','0')
file=OUT/'weather-agent-benchmark-briefing.pptx';prs.save(file)
manifest={'file':str(file.relative_to(ROOT)),'slides':len(prs.slides),'snapshot':stamp,'recorded_attempts':n,'planned_attempts':planned,'native_diagrams':True,'native_chart':True,'raster_assets':['dashboard-comparison.png','dashboard-trace.png']}
(OUT/'briefing-manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
print(json.dumps(manifest,indent=2))
